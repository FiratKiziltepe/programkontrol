"""Sistem Excel'i okuma, PDF'den Excel üretimi ve karşılaştırma testleri.

Örnek sistem Excel'i (samples/system-example.xlsx) Arnavutça 5. sınıfın 4 temasını içerir.
"""
import io

import openpyxl
import pytest

from compare import AYNI, BICIM, EXCEL_VAR, INCELEME, METIN_FARKLI, PDF_VAR, compare, summary, text_status
from excel_export import SYSTEM_HEADERS, build_comparison_excel, build_pdf_excel, system_rows
from excel_import import read_system_excel, strip_placeholder

from conftest import SAMPLES

XL = SAMPLES / "system-example.xlsx"
PDF = SAMPLES / "arnavutca.pdf"


@pytest.fixture(scope="module")
def res():
    if not PDF.exists():
        pytest.skip("samples/arnavutca.pdf yok")
    from validators import analyze_pdf

    return analyze_pdf(str(PDF))


@pytest.fixture(scope="module")
def xl():
    if not XL.exists():
        pytest.skip("samples/system-example.xlsx yok")
    return read_system_excel(str(XL))


@pytest.fixture(scope="module")
def cmp(res, xl):
    return compare(res, xl)


# ---------------------------------------------------------------- okuma


def test_text_status_rules():
    assert text_status("a  b\nc", "a b c\nSayfa(lar)/e-içerik(ler):\n") == AYNI  # sistem eki yok sayılır
    assert text_status("izleme-\nye hazırlık", "izlemeye hazırlık") == BICIM  # satır sonu tiresi
    assert text_status("Beden Eğitimi ve Spor", "Beden Eğitimi Ve Spor") == BICIM  # büyük/küçük harf
    assert text_status("OB5. Kültür\nOkuryazarlığı", "OB5. Kültür Okuryazarlıği") == METIN_FARKLI  # harf farkı
    assert text_status("yansıtır", "yansıtır.") == METIN_FARKLI  # noktalama
    assert strip_placeholder("X\nSayfa(lar)/e-içerik(ler):\n") == "X"


def test_system_excel_structure(xl):
    assert xl.headers == SYSTEM_HEADERS
    assert [t.cell.address for t in xl.themes] == ["A2", "A25", "A48", "A71"]
    assert [len(t.los) for t in xl.themes] == [7, 7, 7, 7]
    # boş Tema/ÖÇ + dolu Süreç Bileşeni satırları önceki öğrenme çıktısına bağlanır
    assert [len(lo.components) for lo in xl.themes[0].los] == [4, 4, 6, 6, 1, 1, 1]
    levels = xl.column_level
    assert levels["Köprü Kurma"] == "theme" and levels["Eğilimler"] == "lo"
    assert levels["Disiplinler Arası İlişkiler"] == "theme"


# ---------------------------------------------------------------- PDF'den Excel


def test_pdf_excel_is_system_format_and_contains_only_pdf_text(res):
    rows, merges = system_rows(res)
    assert len(rows) == sum(max(len(lo.components), 1) for u in res.units for lo in u.learning_outcomes)
    # Metin üretilmez: her hücre PDF'den çıkarılmış bir alanın ham metnidir (kod hücrelerinde satır satır)
    pdf_texts = {res_t for u in res.units for res_t in _texts(u)}
    for r in rows:
        for v in r:
            if not v:
                continue
            for part in [v] if v in pdf_texts else v.split("\n\n"):
                assert part in pdf_texts, part[:80]
    wb = openpyxl.load_workbook(io.BytesIO(build_pdf_excel(res)))
    assert wb.sheetnames == ["Sistem Biçimi", "Tüm Bölümler", "Öğrenme Çıktıları", "Birimler", "Kod Kontrolleri", "Bulgular"]
    assert [c.value for c in wb["Sistem Biçimi"][1]] == SYSTEM_HEADERS


def _texts(u):
    yield u.title.text
    for s in u.sections:
        if s.content:
            yield s.content.text
    for lo in u.learning_outcomes:
        yield lo.header.text
        for c in lo.components:
            yield c.full.text
    for b in u.applications:
        for c in b.used_codes:
            yield c.raw
    for ds in u.declarations.values():
        for d in ds:
            yield d.entry.text


# ---------------------------------------------------------------- karşılaştırma


def test_comparison_directions(cmp, res):
    # Excel yalnızca 5. sınıfı içerir: PDF'deki diğer 12 tema "PDF'de var, Excel'de yok"
    missing = cmp[(cmp["Durum"] == PDF_VAR) & (cmp["Alan / sütun"] == "Tema")]
    assert len(missing) == 12
    assert not missing["Tema (PDF)"].str.startswith("5. SINIF").any()
    # Excel'de olup PDF'de olmayan tek öğe: ARN.5.2.O'nun kod sütunundaki "KB2 Bütünleşik Beceriler (KB2)"
    [row] = cmp[cmp["Durum"] == EXCEL_VAR].to_dict("records")
    assert row["Öğrenme çıktısı"] == "ARN.5.2.O" and row["Öğe"] == "KB2" and row["Excel hücresi"] == "G29"


def test_comparison_text_sections_match(cmp):
    t5 = cmp[cmp["Tema (PDF)"].str.startswith("5. SINIF")]
    for alan in ("Ön Değerlendirme Süreci", "Köprü Kurma", "Tema"):
        assert set(t5[t5["Alan / sütun"] == alan]["Durum"]) == {AYNI}, alan
    lo = t5[t5["Alan / sütun"] == "Öğrenme Çıktısı"]
    assert len(lo) == 28 and set(lo["Durum"]) <= {AYNI, BICIM}
    apps = t5[t5["Alan / sütun"].str.startswith("Öğrenme-Öğretme Uygulamaları")]
    assert len(apps[apps["Öğe"] == "Uygulama metni"]) == 28 and set(apps["Durum"]) <= {AYNI, BICIM}


def test_comparison_real_differences(cmp):
    t5 = cmp[cmp["Tema (PDF)"].str.startswith("5. SINIF")]
    text_diffs = t5[(t5["Durum"] == METIN_FARKLI) & ~t5["Açıklama"].str.contains("ad farklı")]
    # PDF'de "ç) ... yansıtır" (noktasız), sistem Excel'inde "yansıtır."
    assert list(text_diffs["Öğrenme çıktısı"]) == ["ARN.5.4.D"] and list(text_diffs["Öğe"]) == ["ç)"]
    # kod aynı, ad farklı: yalnızca "Okuryazarlıği" yazım farkı. Sistem adlarının sonundaki "Becerisi" eki
    # (kullanıcı kararı) yok sayılır ve not düşülür; başka ad farkı yok sayılmaz.
    names = t5[t5["Açıklama"] == "Kod aynı; ad farklı"]
    assert set(names["Öğe"]) == {"OB5"}
    suffixed = t5[t5["Açıklama"] == 'Adın sonundaki "Becerisi" eki yok sayıldı']
    assert "KB2.2" in set(suffixed["Öğe"]) and set(suffixed["Durum"]) <= {AYNI, BICIM}
    # sistem Excel'inde aynı hücrede tekrarlı kod (D12.1 iki kez)
    dup = t5[t5["Durum"] == INCELEME]
    assert set(dup["Öğe"]) == {"D12.1"} and len(dup) == 3


def test_lo_code_columns_match_application_codes(cmp):
    t5 = cmp[cmp["Tema (PDF)"].str.startswith("5. SINIF")]
    for alan in ("Erdem-Değer-Eylem Çerçevesi", "Eğilimler", "Okuryazarlık Becerileri", "Sosyal-Duygusal Öğrenme Becerileri"):
        part = t5[t5["Alan / sütun"] == alan]
        assert not set(part["Durum"]) & {PDF_VAR, EXCEL_VAR}, alan


def test_comparison_report_excel(res, cmp):
    wb = openpyxl.load_workbook(io.BytesIO(build_comparison_excel(res, cmp, summary(cmp))))
    assert wb.sheetnames == ["Özet", "Karşılaştırma", "PDF Bulguları"]
    assert wb["Karşılaştırma"].max_row == len(cmp) + 1
