"""Tablo karşılaştırma (sistem Excel'i ↔ PDF Excel'i yan yana), metin katmanından/Gemini ile doğrulanmış
tamamlama ve inceleme raporu testleri. Örnek sistem Excel'i Arnavutça 5. sınıfın 4 temasını içerir."""
import io
import json

import docx
import openpyxl
import pytest

import grid_compare as gc
import review_report as rr

from conftest import SAMPLES

PDF = SAMPLES / "arnavutca.pdf"
XL = SAMPLES / "system-example.xlsx"


@pytest.fixture(scope="module")
def base():
    if not (PDF.exists() and XL.exists()):
        pytest.skip("örnek dosyalar yok")
    from excel_import import read_system_excel
    from pdf_extract import extract_spans
    from validators import analyze_pdf

    res = analyze_pdf(str(PDF))
    return res, read_system_excel(str(XL)), extract_spans(res.source)


def _grid(base):
    res, xl, _ = base
    return gc.build_grid(res, xl)


def test_grid_statuses_are_deterministic(base):
    g = _grid(base)
    counts = gc.grid_counts(g)
    assert {k: counts[k] for k in gc.CELL_STATUSES} == {"AYNI": 200, "BİÇİM": 60, "FARKLI": 21, "YALNIZCA_SİSTEM": 0, "YALNIZCA_PDF": 847, "BOŞ": 9}
    # Excel yalnızca 5. sınıfı içerir: 4 tema iki tarafta, 12 tema yalnızca PDF'de
    assert list(g.theme_side.values()).count("both") == 4 and list(g.theme_side.values()).count("pdf") == 12
    # içerik (metin) farkı olarak yalnızca ARN.5.4.D ç) kalır (PDF "yansıtır", sistem "yansıtır.")
    text_diffs = [
        (r.lo_code, r.component)
        for r in g.rows
        if g.theme_side[r.theme_id] == "both" and g.comp_cells[r.id].status == gc.FARKLI
    ]
    assert text_diffs == [("ARN.5.4.D", "ç)")]


def test_code_and_item_lists(base):
    g = _grid(base)
    first = next(r for r in g.rows if r.lo_code == "ARN.5.1.D")
    ok = g.lo_cells[first.lo_id]["Okuryazarlık Becerileri"]
    assert ok.status == gc.FARKLI and ok.note == "Kod aynı, ad farklı: OB5"  # sistem "Okuryazarlıği"
    val = g.lo_cells[first.lo_id]["Erdem-Değer-Eylem Çerçevesi"]
    assert val.status == gc.AYNI and "ad karşılaştırılamadı" in val.note  # D3.1 vb. PDF'de yalnızca kod
    kb = g.lo_cells[first.lo_id]["Beceriler Arası İlişkiler (Alan/Kavramsal Beceriler)"]
    assert kb.status == gc.BICIM and 'Adın sonundaki "Becerisi" eki yok sayıldı' in kb.note  # KB2.12 "Bilgiye/ Veriye" boşluk
    dis = g.theme_cells[first.theme_id]["Disiplinler Arası İlişkiler"]
    assert dis.kind == "items" and dis.status in (gc.AYNI, gc.BICIM)


def test_text_layer_completion_uses_raw_pdf_text(base):
    res, _, doc = base
    g = _grid(base)
    c = g.theme_cells[0]["Köprü Kurma"]
    original = c.pdf
    c.pdf, c.pdf_pages = "", []
    c.recompute()
    assert c.status == gc.SADECE_SISTEM
    assert gc.complete_from_text_layer(g, doc, res.schema_.label_colors) == 1
    assert c.pdf_source == gc.SRC_TEXT_LAYER and c.status == gc.AYNI
    assert gc.compare_key(c.pdf) == gc.compare_key(original)


def test_gemini_text_accepted_only_if_found_in_text_layer(base):
    res, _, doc = base
    g = _grid(base)
    rid = next(r.id for r in g.rows if g.theme_side[r.theme_id] == "both" and g.comp_cells[r.id].status == gc.AYNI)
    real = g.comp_cells[rid]
    true_text = real.pdf
    real.pdf = ""
    real.recompute()
    fake = g.comp_cells[rid + 1]
    fake.pdf = ""
    fake.recompute()
    page = real.pdf_pages[0]
    # Gemini: gerçek metni farklı boşluk/satır kırılımıyla, ikinci alan için PDF'de olmayan bir metin bildirir
    results = {
        f"p{page}": {
            "ok": json.dumps(
                {
                    "fields": [
                        {"id": f"R{rid}|Süreç Bileşeni", "text": " ".join(true_text.split())},
                        {"id": f"R{rid + 1}|Süreç Bileşeni", "text": "Bu cümle PDF'de hiç geçmeyen uydurma bir metindir."},
                    ]
                }
            )
        },
        "p999": {"error": "HTTP 400"},
    }
    stats = gc.apply_gemini(g, doc, res.schema_.label_colors, results, json.loads)
    assert stats == {"verified": 1, "unverified": 1, "errors": 1}
    assert real.pdf_source == gc.SRC_GEMINI and real.pdf == true_text  # PDF'deki ham metin (Gemini'nin değil)
    assert fake.pdf == "" and fake.gemini_unverified.startswith("Bu cümle")  # tabloya alınmadı


def test_gemini_tasks_are_page_based(base):
    res, _, _ = base
    g = _grid(base)
    targets = gc.gemini_targets(g)
    assert targets and all(c.status in (gc.FARKLI, gc.SADECE_SISTEM) for _, c in targets)
    tasks = gc.gemini_tasks(g, res.source, targets[:3], lambda pdf, p: "IMG")
    assert all(t["image"] == "IMG" and t["schema"] == gc.GEMINI_FIELD_SCHEMA and 'id="' in t["text"] for t in tasks)


def test_report_docx_and_xlsx(base):
    g = _grid(base)
    md = rr.python_report(g, "arnavutca.pdf", "system-example.xlsx")
    assert "## 1. Özet" in md and "Farklı: **21**" in md
    # tamamen tek taraflı temalar tek satırda özetlenir
    assert md.count("Tema sistem Excel'inde yok") == 12
    d = docx.Document(io.BytesIO(rr.build_docx(g, md, "arnavutca.pdf", "system-example.xlsx")))
    assert any(p.text == "Ek: Deterministik fark tablosu" for p in d.paragraphs)
    assert len(d.tables[-1].rows) == len(rr.diff_items(g)) + 1
    wb = openpyxl.load_workbook(io.BytesIO(rr.build_grid_xlsx(g)))
    ws = wb.active
    assert ws.max_row == len(g.rows) + 1 and ws.cell(1, 6).value == "Tema — Sistem"
    prompt = rr.gemini_report_prompt(md)
    assert md in prompt and "EKLEME" in rr.GEMINI_REPORT_SYSTEM


def test_payload_is_json_serializable(base):
    g = _grid(base)
    p = gc.grid_payload(g)
    assert len(p["rows"]) == len(g.rows) and json.dumps(p, ensure_ascii=False)


def test_only_becerisi_suffix_is_ignored_in_names():
    # Kullanıcı kararı: yalnızca ad sonundaki "Becerisi" eki yok sayılır; diğer ad farkları fark olarak kalır
    f = gc.name_entry_status
    assert f("KB2.2. Gözlemleme Becerisi", "KB2.2. Gözlemleme") == gc.AYNI
    assert f("KB2.2. Gözlemleme BECERİSİ", "KB2.2. Gözlemleme") == gc.AYNI
    assert f("KB2.4. Çözümleme Becerisi", "KB2.4 Çözümleme") == gc.BICIM  # kod yazımı farklı
    assert f("OB5. Kültür Okuryazarlıği", "OB5. Kültür Okuryazarlığı") == gc.FARKLI
    assert f("KB2.2. Gözlemleme Becerileri", "KB2.2. Gözlemleme") == gc.FARKLI
    assert f("KB2.2. Becerisi Gözlemleme", "KB2.2. Gözlemleme") == gc.FARKLI
    assert f("KB2.2. Gözlemleme Becerisi", "KB2.2. Gözlem") == gc.FARKLI
    assert f("KB2.2. Gözlemleme Becerisi", "KB2.2. Gözlemleme Becerisi") == gc.AYNI
