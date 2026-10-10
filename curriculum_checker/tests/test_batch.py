"""Toplu çalıştırma: dosya adları, elle yapı dosyası, önbellek."""
import shutil

import pytest

import batch
from conftest import SAMPLES


def test_safe_name():
    assert batch.safe_name("2025-BEDEN EĞİTİM PROGRAM-21.07.2025 ") == "2025-BEDEN EĞİTİM PROGRAM-21.07.2025"
    assert batch.safe_name('a:b?c"') == "a_b_c_"
    assert batch.safe_name("...") == "pdf"


def test_overrides_file_roundtrip(tmp_path):
    import openpyxl

    path = tmp_path / batch.OVERRIDES_FILE
    assert batch.ensure_overrides_file(path, ["a.pdf", "b.pdf", "c.pdf"]) is None
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws["B3"], ws["C3"] = "1. ÜNİTE: X", "HMU.11.1.1."
    ws["D4"], ws["E4"] = "yok", 5
    ws["D5"] = "abc"
    wb.save(path)
    batch.ensure_overrides_file(path, ["a.pdf", "b.pdf", "c.pdf", "d.pdf"])  # yeni PDF sona eklenir, girilenler korunur
    ov, err = batch.read_overrides(path)
    assert ov["a.pdf"].unit_title_example == "1. ÜNİTE: X" and ov["a.pdf"].lo_code_example == "HMU.11.1.1."
    assert ov["b.pdf"].structure_pages == [] and ov["b.pdf"].data_start_page == 5
    assert "c.pdf" in err and "c.pdf" not in ov
    names = [r[0] for r in openpyxl.load_workbook(path).active.iter_rows(min_row=3, values_only=True)]
    assert names == ["a.pdf", "b.pdf", "c.pdf", "d.pdf"]


def test_batch_run_and_cache(tmp_path, capsys):
    src = SAMPLES / "temeldinibilgiler.pdf"
    if not src.exists():
        pytest.skip("örnek PDF yok")
    shutil.copy(src, tmp_path / "tdb.pdf")
    assert batch.main([str(tmp_path), "--is-parcacigi", "1"]) == 0
    out = tmp_path / batch.OUT_DIR
    assert (out / "ozet.xlsx").exists() and (out / "tdb" / "tdb_pdf_cikarim.xlsx").exists()
    batch.main([str(tmp_path), "--is-parcacigi", "1"])
    assert "1 tanesi güncel (atlandı), 0 tanesi işlenecek" in capsys.readouterr().out
