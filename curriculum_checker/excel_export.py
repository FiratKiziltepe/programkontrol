"""PDF çıkarım sonucundan tablolar (arayüz) ve indirilebilir Excel üretimi.

Hücrelere yalnızca PDF'den çıkarılmış ham metin yazılır; hiçbir metin üretilmez.
Sayfa 1 ("Sistem Biçimi") müfredat sisteminden indirilen Excel ile aynı sütun ve satır
düzenindedir; böylece iki dosya elle de karşılaştırılabilir.
"""
from __future__ import annotations

import io
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from models import Code, ExtractionResult, Section, SectionKind, Unit
from pdf_extract import norm_label

# Sistem Excel'inin sütun düzeni (sistemin dışa aktarım biçimi; programdan bağımsızdır)
SYSTEM_HEADERS = [
    "Tema",
    "Öğrenme Çıktısı",
    "Süreç Bileşeni",
    "Ön Değerlendirme Süreci",
    "Köprü Kurma",
    "Öğrenme-Öğretme Uygulamaları / Programlar Arası Bileşenlere Yönelik Uygulamalar",
    "Beceriler Arası İlişkiler (Alan/Kavramsal Beceriler)",
    "Erdem-Değer-Eylem Çerçevesi",
    "Eğilimler",
    "Okuryazarlık Becerileri",
    "Sosyal-Duygusal Öğrenme Becerileri",
    "Beceriler Arası İlişkiler (Tema)",
    "Disiplinler Arası İlişkiler",
]
# ÖÇ düzeyindeki kod sütunları: sistem sütunu -> PDF'deki tanım bölümü başlığında aranacak ipuçları.
# Bu sütunlara, o çıktının uygulama bloğunda kullanılan ve ilgili bölümde tanımlı olan kodlar yazılır.
SYSTEM_LO_CODE_COLUMNS = {
    "Beceriler Arası İlişkiler (Alan/Kavramsal Beceriler)": ("kavramsal", "becerilerarasi"),
    "Erdem-Değer-Eylem Çerçevesi": ("degerler",),
    "Eğilimler": ("egilimler",),
    "Okuryazarlık Becerileri": ("okuryazarlik",),
    "Sosyal-Duygusal Öğrenme Becerileri": ("sosyalduygusal",),
}

STATUS_COLORS = {
    # PRD: yeşil aynı, mavi biçim farkı, sarı uyarı / kaynak anomalisi, kırmızı gerçek hata, mor yanlış bölüm
    "AYNI": "C6EFCE",
    "SADECE_BİÇİM_FARKI": "BDD7EE",
    "İNCELEME_GEREKLİ": "FFEB9C",
    "WARNING": "FFEB9C",
    "NEEDS_REVIEW": "FFEB9C",
    "PASS": "C6EFCE",
    "INFO": "EDEDED",  # yalnızca bilgi (ör. süre tablosu bulunamadı)
    "YANLIŞ_BÖLÜM": "D9C3E9",
}
RED = "FFC7CE"


def _xl_safe(v):
    """Excel'e yazılamayan kontrol karakteri (PDF metin katmanındaki görünmez glif) kalmışsa görünür "�"
    ile işaretlenir; hücre sessizce değiştirilmez, dışa aktarma da çökmez."""
    return ILLEGAL_CHARACTERS_RE.sub("�", v) if isinstance(v, str) else v


def _append(ws, row) -> None:
    ws.append([_xl_safe(v) for v in row])


def _ascii_norm(s: str) -> str:
    return (
        norm_label(s)
        .replace("ğ", "g").replace("ş", "s").replace("ç", "c").replace("ö", "o").replace("ü", "u")
    )


def _strip_parens(h: str) -> str:
    return re.sub(r"\([^)]*\)", "", h)


def section_for_header(unit: Unit, header: str) -> Section | None:
    """Sistem sütun başlığına karşılık gelen PDF bölümü (başlık metni eşleştirmesi).

    Tam eşleşme, ya da sütun başlığı bölüm başlığıyla başlıyorsa (ör. "Öğrenme-Öğretme
    Uygulamaları / Programlar Arası ..." -> "Öğrenme-Öğretme Uygulamaları"). Parantez içi yok sayılır.
    """
    hn = norm_label(_strip_parens(header))
    best: Section | None = None
    for s in unit.sections:
        if not s.label_norm or s.kind == SectionKind.GROUP:
            continue
        ln = s.label_norm
        if hn == ln or (hn.startswith(ln) and len(ln) >= 8):
            if best is None or len(ln) > len(best.label_norm):
                best = s
    if best is None and hn.startswith(norm_label("Öğrenme-Öğretme Uygulamaları")):
        # başlığı PDF'de eksik (başlıksız) uygulamalar bölümü
        best = next((s for s in unit.sections if s.kind == SectionKind.APPLICATIONS), None)
    return best


def lo_code_column_codes(unit: Unit, lo_key, header: str) -> list[tuple[Code, str | None]]:
    """Bir ÖÇ'nin uygulama bloğunda kullanılan ve sistem sütununa düşen kodlar: [(kod, PDF'deki tanım metni)].

    Kodun sütunu, onu kapsayan tanım bölümünün başlığından bulunur (D3.1 -> "Değerler" bölümündeki D3).
    Tanım metni yalnızca kodun kendisi (aynı segmentlerle) tanımlıysa döner; değilse None.
    """
    hints = SYSTEM_LO_CODE_COLUMNS.get(header, ())
    block = next((b for b in unit.applications if b.code is not None and b.code.key == lo_key), None)
    if block is None:
        return []
    out: list[tuple[Code, str | None]] = []
    seen = set()
    for c in block.used_codes:
        if c.key in seen:
            continue
        cat = None
        exact = None
        for path, ds in unit.declarations.items():
            for d in ds:
                if d.code.prefix == c.prefix and c.segments[: len(d.code.segments)] == d.code.segments:
                    cat = cat or path
                    if d.code.segments == c.segments:
                        exact, cat = d.entry.text, path
        if cat is None:
            # tanımlı değil: aynı önekin tanımlandığı bölüm
            cat = next((p for p, ds in unit.declarations.items() for d in ds if d.code.prefix == c.prefix), None)
        if cat is not None and any(h in _ascii_norm(cat.split(">")[-1]) for h in hints):
            out.append((c, exact))
            seen.add(c.key)
    # sistem Excel'iyle elle karşılaştırmayı kolaylaştırmak için kod sırasına dizilir (metin değişmez)
    return sorted(out, key=lambda ce: code_sort_key(ce[0]))


def code_sort_key(c: Code):
    return (c.prefix, [int(x) if x.isdigit() else 0 for x in c.segments], c.segments)


# ---------------------------------------------------------------- arayüz tabloları


def _unit_label(u: Unit) -> str:
    return f"{u.context.text + ' | ' if u.context else ''}{u.title.text}".replace("\n", " ")


def units_df(res: ExtractionResult) -> pd.DataFrame:
    rep = {r.unit_id: r for r in res.unit_reports}
    return pd.DataFrame(
        [
            {
                "Birim": u.id,
                "Bağlam": u.context.text if u.context else "",
                "Başlık": u.title.text,
                "Sayfalar": f"{u.pages[0]}-{u.pages[-1]}",
                "Beklenen ÖÇ": rep[u.id].expected_lo,
                "Bulunan ÖÇ": rep[u.id].extracted_lo,
                "Beklenen ders saati": rep[u.id].expected_hours,
                "Bulunan ders saati": rep[u.id].extracted_hours,
                "Durum": rep[u.id].status.value,
            }
            for u in res.units
        ]
    )


def sections_df(res: ExtractionResult, unit_id: str | None = None) -> pd.DataFrame:
    rows = []
    for u in res.units:
        if unit_id and u.id != unit_id:
            continue
        for s in u.sections:
            rows.append(
                {
                    "Birim": u.id,
                    "Tema": _unit_label(u),
                    "Bölüm": s.path or ("(başlıksız: tema açıklaması)" if s.kind == SectionKind.UNIT_DESCRIPTION else "(başlıksız)"),
                    "Tür": s.kind.value,
                    "Metin": s.content.text if s.content else "",
                    "Sayfalar": ", ".join(map(str, s.content.pages)) if s.content else "",
                    "Kaynak span": ", ".join(s.content.source_spans[:3]) + (" …" if s.content and len(s.content.source_spans) > 3 else "") if s.content else "",
                }
            )
    return pd.DataFrame(rows)


def los_df(res: ExtractionResult, unit_id: str | None = None) -> pd.DataFrame:
    rows = []
    for u in res.units:
        if unit_id and u.id != unit_id:
            continue
        blocks = {b.code.key: b for b in u.applications if b.code is not None}
        for lo in u.learning_outcomes:
            b = blocks.get(lo.code.key)
            base = {
                "Birim": u.id,
                "Tema": _unit_label(u),
                "ÖÇ kodu (ham)": lo.code.raw,
                "ÖÇ kodu (normalize)": lo.code.normalized,
                "ÖÇ başlığı": lo.title.text,
                "Kullanılan kodlar (uygulama)": ", ".join(dict.fromkeys(c.normalized for c in b.used_codes)) if b else "",
                "Sayfalar": ", ".join(map(str, lo.title.pages)),
            }
            if not lo.components:
                rows.append({**base, "Bileşen": "", "Bileşen metni": ""})
            for c in lo.components:
                rows.append({**base, "Bileşen": c.marker or "(işaretsiz)", "Bileşen metni": c.text.text})
    return pd.DataFrame(rows)


def findings_df(res: ExtractionResult) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Önem": f.severity.value,
                "Kontrol": f.check,
                "Birim": f.unit_id or "",
                "Açıklama": f.message,
                "Ayrıntı": ", ".join(f"{k}={v}" for k, v in f.details.items() if k not in ("spans", "label_spans"))[:400],
            }
            for f in res.findings
        ]
    )


def code_checks_df(res: ExtractionResult) -> pd.DataFrame:
    rows = []
    for r in res.unit_reports:
        for c in r.code_checks:
            rows.append(
                {
                    "Birim": r.unit_id,
                    "Kategori": c.category,
                    "Tanımlanan": ", ".join(c.declared),
                    "Kullanılan": ", ".join(c.used + c.used_not_declared),
                    "Tanımlı ama kullanılmamış": ", ".join(c.declared_not_used),
                    "Kullanılmış ama tanımlı değil": ", ".join(c.used_not_declared),
                }
            )
        placed = {x for c in r.code_checks for x in c.used_not_declared}
        unplaced = [x for x in r.used_not_declared if x not in placed]
        if unplaced:
            # önek/numarasıyla hiçbir tanım kategorisine yerleşmeyen tanımsız kodlar
            rows.append(
                {
                    "Birim": r.unit_id,
                    "Kategori": "(kategorisi belirlenemeyen)",
                    "Tanımlanan": "",
                    "Kullanılan": ", ".join(unplaced),
                    "Tanımlı ama kullanılmamış": "",
                    "Kullanılmış ama tanımlı değil": ", ".join(unplaced),
                }
            )
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- sistem biçimi


def system_rows(res: ExtractionResult) -> tuple[list[list[str]], list[tuple[int, int, int]]]:
    """Sistem Excel düzeninde satırlar ve birleştirilecek aralıklar [(sütun, ilk satır, son satır)] (1 tabanlı, başlık hariç)."""
    rows: list[list[str]] = []
    merges: list[tuple[int, int, int]] = []
    for u in res.units:
        t_start = len(rows)
        theme_vals = {}
        for h in SYSTEM_HEADERS[3:]:
            if h in SYSTEM_LO_CODE_COLUMNS:
                continue
            sec = section_for_header(u, h)
            theme_vals[h] = sec.content.text if sec and sec.content else ""
        for lo in u.learning_outcomes:
            lo_start = len(rows)
            comps = lo.components or [None]
            for i, c in enumerate(comps):
                row = [""] * len(SYSTEM_HEADERS)
                if len(rows) == t_start:
                    row[0] = u.title.text
                    for h, v in theme_vals.items():
                        row[SYSTEM_HEADERS.index(h)] = v
                if i == 0:
                    row[1] = lo.header.text if lo.header else f"{lo.code.raw}"
                    for h in SYSTEM_LO_CODE_COLUMNS:
                        # girdiler sistemdeki gibi boş satırla ayrılır (bir girdi kendi içinde satır sonu içerebilir)
                        row[SYSTEM_HEADERS.index(h)] = "\n\n".join(
                            exact if exact else c2.raw for c2, exact in lo_code_column_codes(u, lo.code.key, h)
                        )
                row[2] = (c.full.text if c.full else c.text.text) if c else ""
                rows.append(row)
            if len(rows) - lo_start > 1:
                for h in ["Öğrenme Çıktısı", *SYSTEM_LO_CODE_COLUMNS]:
                    merges.append((SYSTEM_HEADERS.index(h) + 1, lo_start + 2, len(rows) + 1))
        if len(rows) - t_start > 1:
            for h in ["Tema", *theme_vals]:
                merges.append((SYSTEM_HEADERS.index(h) + 1, t_start + 2, len(rows) + 1))
    return rows, merges


def _write_df(ws, df: pd.DataFrame, status_col: str | None = None):
    _append(ws, list(df.columns))
    for c in ws[1]:
        c.font = Font(bold=True)
    for rec in df.itertuples(index=False):
        _append(ws, [("" if v is None else v) for v in rec])
    if status_col and status_col in df.columns:
        idx = list(df.columns).index(status_col) + 1
        for r in range(2, ws.max_row + 1):
            v = str(ws.cell(r, idx).value)
            fill = PatternFill("solid", fgColor=STATUS_COLORS.get(v, RED))
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).fill = fill
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 40
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def build_pdf_excel(res: ExtractionResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sistem Biçimi"
    _append(ws, SYSTEM_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    rows, merges = system_rows(res)
    for r in rows:
        _append(ws, r)
    for col, r1, r2 in merges:
        ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 45
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _write_df(wb.create_sheet("Tüm Bölümler"), sections_df(res))
    _write_df(wb.create_sheet("Öğrenme Çıktıları"), los_df(res))
    _write_df(wb.create_sheet("Birimler"), units_df(res), "Durum")
    _write_df(wb.create_sheet("Kod Kontrolleri"), code_checks_df(res))
    fdf = findings_df(res)
    _write_df(wb.create_sheet("Bulgular"), fdf, "Önem")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_comparison_excel(res: ExtractionResult, comparison: pd.DataFrame, summary: pd.DataFrame) -> bytes:
    """Karşılaştırma raporu: özet + renkli sonuç tablosu (PRD renkleri) + PDF bulguları."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    _append(ws, ["Kaynak PDF", res.source.split("\\")[-1].split("/")[-1]])
    _append(ws, ["PDF genel durum", res.status.value])
    _append(ws, ["Tema/ünite (beklenen / bulunan)", f"{res.expected_unit_count} / {len(res.units)}"])
    _append(ws, ["Öğrenme çıktısı (beklenen / çıkarılan)", f"{res.expected_lo_total} / {res.extracted_lo_total}"])
    _append(ws, [])
    _write_df_into(ws, summary, "Durum")
    _write_df(wb.create_sheet("Karşılaştırma"), comparison, "Durum")
    _write_df(wb.create_sheet("PDF Bulguları"), findings_df(res), "Önem")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_df_into(ws, df: pd.DataFrame, status_col: str):
    start = ws.max_row + 1
    _append(ws, list(df.columns))
    for rec in df.itertuples(index=False):
        _append(ws, list(rec))
    idx = list(df.columns).index(status_col) + 1
    for r in range(start + 1, ws.max_row + 1):
        ws.cell(r, idx).fill = PatternFill("solid", fgColor=STATUS_COLORS.get(str(ws.cell(r, idx).value), RED))
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 30
