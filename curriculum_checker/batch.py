"""Klasördeki tüm öğretim programı PDF'lerini toplu analiz eder (Streamlit gerekmez).

Çalıştırma:
    python batch.py "C:\\Users\\...\\programlar"            # çıktılar: <klasör>\\_cikti
    python batch.py "<klasör>" --yeniden                     # tüm PDF'leri baştan işle
    python batch.py "<klasör>" --saat-yoksay                 # ders saati uyuşmazlığını göz ardı et

Her PDF için _cikti\\<ad>\\ altına:
  - <ad>_pdf_cikarim.xlsx  : arayüzdeki "PDF'den üretilen Excel" ile aynı
  - <ad>_isaretli.pdf      : Hata + İnceleme gerekli bulguları işaretli PDF (arayüz varsayılanı)
  - ozet.json              : özet satırı (PDF, elle yapı bilgisi ve kod değişmediyse PDF yeniden işlenmez)
_cikti\\ozet.xlsx: her PDF bir satır, otomatik tanıma sorunlarına göre gruplar, bulgu türleri.
<klasör>\\elle_yapi.xlsx: otomatik tanınamayan PDF'ler için elle yapı bilgileri (arayüzdeki formla aynı alanlar).
Dosya yoksa oluşturulur; yeni PDF'ler sona eklenir, girilen değerlere dokunulmaz.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import time
import traceback
from collections import Counter
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

OUT_DIR = "_cikti"
OVERRIDES_FILE = "elle_yapi.xlsx"
OV_HEADERS = ["Dosya", "Örnek tema/ünite başlığı", "Örnek ÖÇ kodu", "Yapı sayfaları", "İlk tema sayfası", "Not"]
OV_HELP = [
    "PDF dosya adı (değiştirmeyin)",
    "PDF'de yazdığı gibi, ör. 1. ÜNİTE: DİN HİZMETLERİ (boş: otomatik)",
    "PDF'de yazdığı gibi, ör. HMU.11.1.1. (boş: otomatik)",
    "ör. 12-13 ya da 12, 13; yapı sayfası yoksa: yok (boş: otomatik)",
    "PDF sayfa numarası (boş: otomatik)",
    "Serbest not (araç okumaz)",
]

# Otomatik tanımanın (çıkarımın) şüpheli olduğunu gösteren bulgular -> özet Excel'deki sorun grubu.
# İçerik hataları (ör. tanımlanmamış kod) burada YOK: onlar PDF'nin kendi hatasıdır, tanıma sorunu değildir.
RECOGNITION_ISSUES = {
    "STRUCTURE_PAGE_NOT_FOUND": "Yapı sayfası bulunamadı",
    "EXPECTED_TABLE_NOT_FOUND": "Süre tablosu bulunamadı",
    "EXPECTED_TABLE_UNMATCHED": "Süre tablosu temalarla eşleşmedi",
    "EXPECTED_TABLE_AMBIGUOUS": "Süre tablosu temalarla eşleşmedi",
    "EXPECTED_ROW_AMBIGUOUS": "Süre tablosu temalarla eşleşmedi",
    "EXPECTED_TOTAL_UNKNOWN": "Süre tablosu temalarla eşleşmedi",
    "EXPECTED_COUNT_UNREADABLE": "Süre tablosu okunamadı",
    "UNIT_COUNT_MISMATCH": "Tema sayısı tutmuyor",
    "MISSING_UNIT": "Tema sayısı tutmuyor",
    "LO_COUNT_MISMATCH": "ÖÇ sayısı tutmuyor",
    "LO_TOTAL_MISMATCH": "ÖÇ sayısı tutmuyor",
    "LO_SECTION_NOT_FOUND": "ÖÇ bölümü bulunamadı",
    "MISSING_SECTION": "Bölüm başlığı bulunamadı",
    "APPLICATIONS_LABEL_MISSING": "Bölüm başlığı bulunamadı",
    "UNKNOWN_LABEL": "Başlık yapısı farklı",
    "ORPHAN_LABEL": "Başlık yapısı farklı",
    "GROUP_STYLE_VARIANT": "Başlık yapısı farklı",
    "GROUP_HAS_CONTENT": "Başlık yapısı farklı",
    "UNASSIGNED_SPANS": "Metin bölümlere tam dağıtılamadı",
    "DOUBLE_ASSIGNED_SPANS": "Metin bölümlere tam dağıtılamadı",
    "TEXT_NOT_IN_SOURCE": "Metin bölümlere tam dağıtılamadı",
    "NO_SOURCE_SPANS": "Metin bölümlere tam dağıtılamadı",
    "UNPARSED_TEXT": "Metin bölümlere tam dağıtılamadı",
    "UNATTACHED_APPLICATION_TEXT": "Metin bölümlere tam dağıtılamadı",
}
SCHEMA_FAIL = "Yapı tanınamadı (elle yapı gerekli)"
CRASH = "Analiz hatası"
NO_UNITS = "Tema bulunamadı"


def safe_name(stem: str, n: int = 80) -> str:
    """Windows'ta klasör/dosya adı olabilecek kısa ad (sonda boşluk/nokta olamaz)."""
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", stem).strip(" .")
    return (s[:n].rstrip(" .") or "pdf")


def code_hash() -> str:
    h = hashlib.sha1()
    for p in sorted(HERE.glob("*.py")):
        h.update(p.read_bytes())
    return h.hexdigest()[:12]


# ---------------------------------------------------------------- elle yapı dosyası


def _cell(v) -> str:
    return "" if v is None else str(v).strip()


def read_overrides(path: Path) -> tuple[dict, dict[str, str]]:
    """{dosya adı: SchemaOverrides}, {dosya adı: hata}. Yanlış biçimli değer o PDF için hata olarak raporlanır."""
    from manual_schema import _parse_pages
    from models import SchemaOverrides

    out, errors = {}, {}
    if not path.exists():
        return out, errors
    import openpyxl

    ws = openpyxl.load_workbook(path, read_only=True).active
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or not _cell(row[0]):
            continue
        name, title, code, pages, start = (_cell(x) for x in (list(row) + [None] * 5)[:5])
        if not (title or code or pages or start):
            continue
        structure = None
        if pages:
            structure = [] if pages.lower() in ("yok", "-", "0") else _parse_pages(pages)
            if structure is None:
                errors[name] = f"Yapı sayfaları okunamadı: {pages!r} (ör. 12-13 ya da yok)"
                continue
        if start and not re.fullmatch(r"\d+(\.0)?", start):
            errors[name] = f"İlk tema sayfası sayı değil: {start!r}"
            continue
        out[name] = SchemaOverrides(
            unit_title_example=title or None,
            lo_code_example=code or None,
            structure_pages=structure,
            data_start_page=int(float(start)) if start else None,
        )
    return out, errors


def ensure_overrides_file(path: Path, names: list[str]) -> str | None:
    """Dosyayı oluşturur ya da eksik PDF adlarını sona ekler; girilmiş değerlere dokunmaz."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    try:
        if path.exists():
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            have = {_cell(r[0]) for r in ws.iter_rows(min_row=3, values_only=True) if r and r[0]}
            new = [n for n in names if n not in have]
            if not new:
                return None
            for n in new:
                ws.append([n])
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Elle yapı"
            ws.append(OV_HEADERS)
            ws.append(OV_HELP)
            for c in ws[1]:
                c.font = Font(bold=True)
            for c in ws[2]:
                c.font = Font(italic=True, color="666666")
                c.fill = PatternFill("solid", fgColor="F2F2F2")
            for n in names:
                ws.append([n])
            for col, w in zip("ABCDEF", (60, 45, 22, 26, 18, 40)):
                ws.column_dimensions[col].width = w
            ws.freeze_panes = "B3"
        wb.save(path)
    except PermissionError:
        return f"{path.name} açık olduğu için güncellenemedi (Excel'de kapatıp yeniden çalıştırın)."
    return None


# ---------------------------------------------------------------- tek PDF


def process_one(pdf: str, out_dir: str, ov_json: str, ignore_hours: bool, signature: str) -> dict:
    """Bir PDF'i analiz eder, Excel + işaretli PDF yazar, özet satırını döndürür (ayrı süreçte çalışır)."""
    from excel_export import build_pdf_excel
    from labels_tr import CHECK_TR, SEVERITY_TR, STATUS_TR
    from models import SchemaOverrides
    from pdf_annotate import build_annotated_pdf
    from pdf_extract import extract_spans
    from pdf_locate import Locator
    from schema_detect import SchemaError
    from validators import analyze_pdf, ignore_checks

    t0 = time.time()
    p, od = Path(pdf), Path(out_dir)
    od.mkdir(parents=True, exist_ok=True)
    base = od.name
    row: dict = {"Dosya": p.name, "Elle yapı": "evet" if ov_json else "", "_imza": signature, "_klasor": od.name}
    ov = SchemaOverrides.model_validate_json(ov_json) if ov_json else None
    try:
        res = analyze_pdf(str(p), ov)
    except SchemaError as e:
        row.update({"Durum": "Yapı tanınamadı", "Tanıma sorunları": [SCHEMA_FAIL], "Hata mesajı": str(e)})
    except Exception as e:
        row.update({"Durum": "Analiz hatası", "Tanıma sorunları": [CRASH], "Hata mesajı": f"{type(e).__name__}: {e}", "_iz": traceback.format_exc()[-3000:]})
    else:
        if ignore_hours:
            res = ignore_checks(res, {"HOURS_MISMATCH"})
        sev = Counter(f.severity.value for f in res.findings)
        serious = [f for f in res.findings if f.severity.value in ("FAIL", "NEEDS_REVIEW")]
        issues = sorted({RECOGNITION_ISSUES[f.check] for f in res.findings if f.check in RECOGNITION_ISSUES})
        if not res.units:
            issues.insert(0, NO_UNITS)
        s = res.schema_
        top = Counter(f.check for f in serious).most_common(6)
        row.update(
            {
                "Durum": STATUS_TR.get(res.status.value, res.status.value),
                "Tanıma sorunları": issues,
                "Tema (bulunan)": len(res.units),
                "Tema (beklenen)": res.expected_unit_count if res.expected_lo_total is not None else None,
                "ÖÇ (bulunan)": res.extracted_lo_total,
                "ÖÇ (beklenen)": res.expected_lo_total,
                SEVERITY_TR["FAIL"]: sev.get("FAIL", 0),
                SEVERITY_TR["NEEDS_REVIEW"]: sev.get("NEEDS_REVIEW", 0),
                SEVERITY_TR["WARNING"]: sev.get("WARNING", 0),
                "Sık bulgular": "; ".join(f"{CHECK_TR.get(c, c)} ({n})" for c, n in top),
                "Anahtar kelime": s.unit_keyword,
                "ÖÇ kodu deseni": f"{s.lo_prefix}." + ".".join("N" if t == "n" else "X" for t in s.lo_segment_types),
                "Yapı sayfaları": ", ".join(map(str, s.structure_pages)) or "yok",
                "İlk tema sayfası": s.data_start_page,
                "Örnek tema başlığı": res.units[0].title.text.replace("\n", " ") if res.units else "",
                "Sayfa sayısı": res.page_count,
                "_bulgular": [f.check for f in serious],
            }
        )
        try:
            (od / f"{base}_pdf_cikarim.xlsx").write_bytes(build_pdf_excel(res))
            row["Excel"] = f"{base}_pdf_cikarim.xlsx"
            if serious:
                L = Locator(res, extract_spans(res.source))
                (od / f"{base}_isaretli.pdf").write_bytes(build_annotated_pdf(res, serious, L))
                row["İşaretli PDF"] = f"{base}_isaretli.pdf"
        except Exception as e:
            row["Hata mesajı"] = f"Çıktı yazılamadı: {type(e).__name__}: {e}"
            row["_iz"] = traceback.format_exc()[-3000:]
    row["Süre (sn)"] = round(time.time() - t0, 1)
    (od / "ozet.json").write_text(json.dumps(row, ensure_ascii=False, indent=1), encoding="utf-8")
    return row


# ---------------------------------------------------------------- özet Excel

SUMMARY_COLS = [
    "Dosya", "Durum", "Tanıma sorunları", "Tema (bulunan)", "Tema (beklenen)", "ÖÇ (bulunan)", "ÖÇ (beklenen)",
    "Hata", "İnceleme gerekli", "Uyarı", "Sık bulgular", "Elle yapı", "Anahtar kelime", "ÖÇ kodu deseni",
    "Yapı sayfaları", "İlk tema sayfası", "Örnek tema başlığı", "Sayfa sayısı", "Excel", "İşaretli PDF", "Hata mesajı", "Süre (sn)",
]


def write_summary(rows: list[dict], out: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    from labels_tr import CHECK_TR

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PDF'ler"
    ws.append(SUMMARY_COLS)
    fills = {"Geçti": "C6EFCE", "Uyarılarla geçti": "FFEB9C", "İnceleme gerekli": "FFEB9C", "Hata": "FFC7CE"}
    rows = sorted(rows, key=lambda r: (not r.get("Tanıma sorunları"), r["Dosya"].lower()))
    for r in rows:
        vals = []
        for c in SUMMARY_COLS:
            v = r.get(c)
            vals.append("; ".join(v) if isinstance(v, list) else v)
        ws.append(vals)
        i = ws.max_row
        for c in ("Excel", "İşaretli PDF"):
            if r.get(c):
                cell = ws.cell(i, SUMMARY_COLS.index(c) + 1)
                cell.hyperlink = f"{r['_klasor']}/{r[c]}"
                cell.font = Font(color="0563C1", underline="single")
        if r.get("Tanıma sorunları"):
            ws.cell(i, 3).fill = PatternFill("solid", fgColor="F8CBAD")
        ws.cell(i, 2).fill = PatternFill("solid", fgColor=fills.get(r.get("Durum"), "D9D9D9"))
    widths = {"A": 50, "B": 18, "C": 45, "K": 70, "Q": 45, "U": 50}
    for col in ws.iter_cols(1, len(SUMMARY_COLS)):
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = widths.get(letter, 14)
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    # Sorun grupları: hangi tanıma sorunu kaç PDF'de (kod düzeltme önceliği)
    groups: dict[str, list[str]] = {}
    for r in rows:
        for s in r.get("Tanıma sorunları") or []:
            groups.setdefault(s, []).append(r["Dosya"])
    g = wb.create_sheet("Sorun grupları")
    g.append(["Tanıma sorunu", "PDF sayısı", "PDF'ler"])
    for s, names in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        g.append([s, len(names), "\n".join(names)])
    g.append([])
    g.append(["Sorunsuz tanınan PDF", sum(1 for r in rows if not r.get("Tanıma sorunları"))])

    # Bulgu türleri (Hata + İnceleme gerekli): tüm PDF'lerde
    per_pdf: dict[str, set[str]] = {}
    total: Counter = Counter()
    for r in rows:
        for c in r.get("_bulgular") or []:
            per_pdf.setdefault(c, set()).add(r["Dosya"])
            total[c] += 1
    b = wb.create_sheet("Bulgu türleri")
    b.append(["Kontrol", "Kod", "Tanıma sorunu mu?", "PDF sayısı", "Toplam bulgu"])
    for c, n in total.most_common():
        b.append([CHECK_TR.get(c, c), c, RECOGNITION_ISSUES.get(c, ""), len(per_pdf[c]), n])

    for sheet, w in ((g, (40, 12, 90)), (b, (45, 30, 35, 12, 14))):
        for i, width in enumerate(w):
            sheet.column_dimensions[chr(65 + i)].width = width
        for c in sheet[1]:
            c.font = Font(bold=True)
        for col in sheet.iter_cols():
            for c in col:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    path = out / "ozet.xlsx"
    try:
        wb.save(path)
    except PermissionError:
        path = out / f"ozet_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(path)
    return path


# ---------------------------------------------------------------- akış


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Klasördeki öğretim programı PDF'lerini toplu analiz eder.")
    ap.add_argument("klasor", nargs="?", default=r"C:\Users\firat\Desktop\programlar")
    ap.add_argument("--yeniden", action="store_true", help="Daha önce işlenmiş PDF'leri de baştan işle")
    ap.add_argument("--saat-yoksay", action="store_true", help="Ders saati uyuşmazlığını göz ardı et")
    ap.add_argument("--is-parcacigi", type=int, default=max(1, (os.cpu_count() or 2) - 1), help="Aynı anda işlenecek PDF sayısı")
    a = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # Türkçe dosya adları her konsolda yazılsın

    root = Path(a.klasor)
    out = root / OUT_DIR
    out.mkdir(exist_ok=True)
    pdfs = sorted((p for p in root.rglob("*") if p.suffix.lower() == ".pdf" and OUT_DIR not in p.relative_to(root).parts), key=lambda p: str(p).lower())
    if not pdfs:
        print(f"{root} içinde PDF yok.")
        return 1
    names = [str(p.relative_to(root)) for p in pdfs]
    ov_path = root / OVERRIDES_FILE
    msg = ensure_overrides_file(ov_path, names)
    if msg:
        print("UYARI:", msg)
    overrides, ov_errors = read_overrides(ov_path)
    chash = code_hash()

    # aynı kısa ada düşen PDF'ler için klasör adları tekilleştirilir
    folders, used = {}, Counter()
    for p in pdfs:
        n = safe_name(p.stem)
        used[n] += 1
        folders[p] = n if used[n] == 1 else f"{n}_{used[n]}"

    rows, todo = [], []
    for p, name in zip(pdfs, names):
        ov = overrides.get(name)
        ov_json = ov.model_dump_json() if ov is not None and not ov.is_empty() else ""
        st = p.stat()
        sig = hashlib.sha1(f"{st.st_size}|{st.st_mtime_ns}|{ov_json}|{a.saat_yoksay}|{chash}".encode()).hexdigest()[:16]
        od = out / folders[p]
        cached = od / "ozet.json"
        if not a.yeniden and cached.exists():
            try:
                row = json.loads(cached.read_text(encoding="utf-8"))
                if row.get("_imza") == sig:
                    row["Dosya"] = name
                    rows.append(row)
                    continue
            except (OSError, ValueError):
                pass
        todo.append((p, name, od, ov_json, sig))

    print(f"{len(pdfs)} PDF; {len(rows)} tanesi güncel (atlandı), {len(todo)} tanesi işlenecek ({a.is_parcacigi} paralel).")
    t0 = time.time()
    if todo:
        with ProcessPoolExecutor(max_workers=a.is_parcacigi) as ex:
            futs = {ex.submit(process_one, str(p), str(od), ov_json, a.saat_yoksay, sig): (p, name, od) for p, name, od, ov_json, sig in todo}
            for k, fut in enumerate(as_completed(futs), 1):
                p, name, od = futs[fut]
                try:
                    row = fut.result()
                except Exception as e:  # süreç çöktü (bellek vb.)
                    row = {"Dosya": name, "Durum": "Analiz hatası", "Tanıma sorunları": [CRASH], "Hata mesajı": f"{type(e).__name__}: {e}", "_klasor": od.name}
                row["Dosya"] = name
                rows.append(row)
                issues = "; ".join(row.get("Tanıma sorunları") or []) or "sorunsuz tanındı"
                print(f"[{k}/{len(todo)}] {name} — {row.get('Durum')} — {issues} ({row.get('Süre (sn)', '?')} sn)", flush=True)
    for name, err in ov_errors.items():
        r = next((r for r in rows if r["Dosya"] == name), None)
        if r is not None:
            r["Hata mesajı"] = ((r.get("Hata mesajı") or "") + f" | elle_yapi.xlsx: {err}").strip(" |")
    path = write_summary(rows, out)
    bad = sum(1 for r in rows if r.get("Tanıma sorunları"))
    print(f"\nBitti ({round(time.time() - t0)} sn). {len(rows) - bad} PDF sorunsuz tanındı, {bad} PDF'de tanıma sorunu var.")
    print(f"Özet: {path}")
    if ov_errors:
        print(f"UYARI: elle_yapi.xlsx'te {len(ov_errors)} satır okunamadı (özet 'Hata mesajı' sütununda).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
