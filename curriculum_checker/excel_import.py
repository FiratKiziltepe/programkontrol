"""Müfredat sisteminden indirilen Excel'in okunması.

- Her değer için sayfa, hücre adresi ve ham hücre değeri korunur.
- Birleştirilmiş hücrelerde değer sol üst hücrededir; boş devam satırları doğru üst kayda bağlanır
  (Tema ve Öğrenme Çıktısı boş, Süreç Bileşeni dolu satır -> önceki öğrenme çıktısının devamı).
- Sütunların düzeyi (tema / öğrenme çıktısı / süreç bileşeni) sütun adından değil, verinin
  hangi satırlarda dolu olduğundan öğrenilir.
- Sistemin her hücreye eklediği "Sayfa(lar)/e-içerik(ler):" satırı karşılaştırmada yok sayılır;
  ham değer olduğu gibi saklanır.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl

from pdf_extract import norm_label

# Sistemin her hücreye eklediği alan etiketi (kullanıcı kararı: karşılaştırmada yok sayılır)
SYSTEM_PLACEHOLDER_RE = re.compile(r"Sayfa\(lar\)\s*/\s*e-içerik\(ler\)\s*:")


def strip_placeholder(value: str | None) -> str:
    """Karşılaştırma için: sistem ekini çıkarır, baştaki/sondaki boşlukları kırpar."""
    if value is None:
        return ""
    return SYSTEM_PLACEHOLDER_RE.sub("", str(value)).strip()


@dataclass
class Cell:
    sheet: str
    address: str
    raw: str | None
    row: int
    col: int

    @property
    def text(self) -> str:
        return strip_placeholder(self.raw)


@dataclass
class ExcelLO:
    cell: Cell  # "Öğrenme Çıktısı" hücresi
    components: list[Cell] = field(default_factory=list)
    columns: dict[str, Cell] = field(default_factory=dict)  # ÖÇ düzeyindeki diğer sütunlar


@dataclass
class ExcelTheme:
    cell: Cell  # "Tema" hücresi
    los: list[ExcelLO] = field(default_factory=list)
    columns: dict[str, Cell] = field(default_factory=dict)  # tema düzeyindeki sütunlar


@dataclass
class SystemExcel:
    path: str
    sheet: str
    headers: list[str]
    theme_col: int
    lo_col: int
    component_col: int
    column_level: dict[str, str]  # başlık -> "theme" | "lo" | "component" | "empty"
    themes: list[ExcelTheme]


def _find_col(headers: list[str], *keys: str) -> int | None:
    for i, h in enumerate(headers):
        if norm_label(h or "") in {norm_label(k) for k in keys}:
            return i + 1
    return None


def read_system_excel(path: str) -> SystemExcel:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    theme_col = _find_col(headers, "Tema", "Ünite", "Öğrenme Alanı")
    lo_col = _find_col(headers, "Öğrenme Çıktısı")
    comp_col = _find_col(headers, "Süreç Bileşeni", "Süreç Bileşenleri")
    if not (theme_col and lo_col and comp_col):
        raise ValueError(f"Beklenen sütunlar bulunamadı (Tema / Öğrenme Çıktısı / Süreç Bileşeni): {headers}")

    def cell(r: int, c: int) -> Cell:
        v = ws.cell(r, c).value
        return Cell(ws.title, ws.cell(r, c).coordinate, None if v is None else str(v), r, c)

    def filled(r: int, c: int) -> bool:
        v = ws.cell(r, c).value
        return v is not None and str(v).strip() != ""

    # Sütun düzeyi: dolu hücreler yalnızca tema satırlarındaysa tema, ÖÇ satırlarındaysa ÖÇ düzeyi
    theme_rows = {r for r in range(2, ws.max_row + 1) if filled(r, theme_col)}
    lo_rows = {r for r in range(2, ws.max_row + 1) if filled(r, lo_col)}
    level: dict[str, str] = {}
    for c, h in enumerate(headers, start=1):
        if c in (theme_col, lo_col, comp_col):
            continue
        rows = {r for r in range(2, ws.max_row + 1) if filled(r, c)}
        if not rows:
            level[h] = "empty"
        elif rows <= theme_rows:
            level[h] = "theme"
        elif rows <= lo_rows:
            level[h] = "lo"
        else:
            level[h] = "component"

    themes: list[ExcelTheme] = []
    cur_t: ExcelTheme | None = None
    cur_lo: ExcelLO | None = None
    for r in range(2, ws.max_row + 1):
        if filled(r, theme_col):
            cur_t = ExcelTheme(cell=cell(r, theme_col))
            for c, h in enumerate(headers, start=1):
                if level.get(h) == "theme" and filled(r, c):
                    cur_t.columns[h] = cell(r, c)
            themes.append(cur_t)
            cur_lo = None
        if cur_t is None:
            continue
        if filled(r, lo_col):
            cur_lo = ExcelLO(cell=cell(r, lo_col))
            for c, h in enumerate(headers, start=1):
                if level.get(h) == "lo" and filled(r, c):
                    cur_lo.columns[h] = cell(r, c)
            cur_t.los.append(cur_lo)
        if filled(r, comp_col) and cur_lo is not None:
            # boş Tema/ÖÇ ile dolu Süreç Bileşeni: önceki öğrenme çıktısının devamı
            cur_lo.components.append(cell(r, comp_col))
    return SystemExcel(path, ws.title, headers, theme_col, lo_col, comp_col, level, themes)


# ---------------------------------------------------------------- hücre içeriği ayrıştırma

CODE_LINE_RE = re.compile(r"^\s*([^\W\d_]{1,5}\d+(?:\.\d+)*)\.?\s*(.*)$")


def code_items(cell: Cell) -> list[tuple[str, str, str]]:
    """Kod listesi hücresi: [(ham_satır, kod, ad)]. Sistem eki satırları atlanır.

    ör. "KB2.2. Gözlemleme Becerisi\\nSayfa(lar)/e-içerik(ler):\\n\\nKB2.5. ..." """
    out = []
    for line in cell.text.split("\n"):
        line = SYSTEM_PLACEHOLDER_RE.sub("", line).strip()
        if not line:
            continue
        m = CODE_LINE_RE.match(line)
        if m:
            out.append((line, m.group(1), m.group(2).strip()))
        else:
            out.append((line, "", line))
    return out


def list_items(cell: Cell) -> list[str]:
    """Kodsuz liste hücresi (ör. Disiplinler Arası İlişkiler): satır başına bir öğe."""
    return [x.strip() for x in cell.text.split("\n") if x.strip()]


# ---------------------------------------------------------------- birden çok Excel


def _tag_cells(theme: ExcelTheme, tag: str) -> None:
    """Birden çok dosyada hücre adresine dosya adı eklenir ("5.sinif.xlsx!A2"); değerler değişmez."""
    cells = [theme.cell, *theme.columns.values()]
    for lo in theme.los:
        cells += [lo.cell, *lo.components, *lo.columns.values()]
    for c in cells:
        c.address = f"{tag}!{c.address}"


def merge_system_excels(parts: list[tuple[str, SystemExcel]]) -> SystemExcel:
    """Sınıf sınıf indirilen sistem Excel'lerini tek listede birleştirir: [(dosya adı, SystemExcel)].
    Temalar dosya sırasıyla eklenir. Sütun düzeyleri birleştirilir; bir sütun dosyalarda farklı düzeyde ise
    ilk dosyadaki geçerlidir. Tek dosyada hiçbir şey değişmez."""
    if len(parts) == 1:
        return parts[0][1]
    first = parts[0][1]
    headers = list(first.headers)
    level = dict(first.column_level)
    themes: list[ExcelTheme] = []
    for name, x in parts:
        for h in x.headers:
            if h not in headers:
                headers.append(h)
        for h, lv in x.column_level.items():
            if level.get(h) in (None, "empty"):
                level[h] = lv
        for t in x.themes:
            _tag_cells(t, name)
            themes.append(t)
    return SystemExcel(
        " + ".join(n for n, _ in parts), first.sheet, headers, first.theme_col, first.lo_col, first.component_col, level, themes
    )
