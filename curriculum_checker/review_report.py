"""İnceleme uzmanları için rapor: nihai (tamamlanmış) karşılaştırma tablosundaki farklardan deterministik olarak
Python ile üretilir. "Gemini ile raporlaştır" isteğe bağlıdır: Gemini yalnızca bu raporu daha okunaklı yazar;
yeni olgu eklemesi, sayıları/kodları değiştirmesi istenmez. Word çıktısında deterministik fark tablosu her
durumda ek olarak yer alır.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from grid_compare import (
    AYNI,
    BICIM,
    BOS,
    CELL_STATUSES,
    DIFF_STATUSES,
    FARKLI,
    SADECE_PDF,
    SADECE_SISTEM,
    SRC_GEMINI,
    SRC_TEXT_LAYER,
    Grid,
    GridCell,
    grid_counts,
)

STATUS_TR = {
    AYNI: "Aynı",
    BICIM: "Biçim farkı",
    FARKLI: "Farklı",
    SADECE_SISTEM: "Yalnızca sistemde",
    SADECE_PDF: "Yalnızca PDF'de",
    BOS: "İki taraf boş",
}


def _short(t: str, n: int = 220) -> str:
    t = re.sub(r"\s+", " ", t or "").strip()
    return t if len(t) <= n else t[: n - 1] + "…"


def diff_items(grid: Grid) -> list[dict]:
    """Nihai tablodaki fark hücreleri (her hücre bir kez), tema sırasıyla."""
    items = []
    seen = set()
    one_sided: set = set()
    for r in grid.rows:
        # Tamamen tek tarafta olan tema / öğrenme çıktısı tek satırla raporlanır (her hücresi ayrı fark değildir)
        side = grid.theme_side.get(r.theme_id, "both")
        lo_side = grid.lo_side.get(r.lo_id, "both") if r.lo_id is not None else "both"
        scope = ("T", r.theme_id) if side != "both" else (("L", r.lo_id) if lo_side != "both" else None)
        if scope is not None:
            if scope not in one_sided:
                one_sided.add(scope)
                s = side if scope[0] == "T" else lo_side
                cell = grid.theme_cells[r.theme_id].get("Tema") if scope[0] == "T" else grid.lo_cells[r.lo_id].get("Öğrenme Çıktısı")
                items.append(
                    {
                        "tema": r.theme_label,
                        "oc": r.lo_code if scope[0] == "L" else "",
                        "bilesen": "",
                        "alan": "Tema (tümü)" if scope[0] == "T" else "Öğrenme çıktısı (tümü)",
                        "durum": SADECE_PDF if s == "pdf" else SADECE_SISTEM,
                        "sistem": cell.sys if cell else "",
                        "pdf": cell.pdf if cell else "",
                        "not": ("Tema" if scope[0] == "T" else "Öğrenme çıktısı") + (" sistem Excel'inde yok" if s == "pdf" else " PDF'de yok"),
                        "kaynak": "",
                        "gemini_dogrulanamayan": "",
                        "excel": cell.sys_address if cell else "",
                        "sayfa": ", ".join(map(str, cell.pdf_pages)) if cell else "",
                    }
                )
            continue
        for h, c in grid.cells_of_row(r).items():
            if c.status not in DIFF_STATUSES and not c.gemini_unverified:
                continue
            if id(c) in seen:
                continue
            seen.add(id(c))
            items.append(
                {
                    "tema": r.theme_label,
                    "oc": r.lo_code if c.level != "theme" else "",
                    "bilesen": r.component if c.level == "component" else "",
                    "alan": h,
                    "durum": c.status,
                    "sistem": c.sys,
                    "pdf": c.pdf,
                    "not": c.note,
                    "kaynak": c.pdf_source,
                    "gemini_dogrulanamayan": c.gemini_unverified,
                    "excel": c.sys_address,
                    "sayfa": ", ".join(map(str, c.pdf_pages)),
                }
            )
    return items


def python_report(grid: Grid, pdf_name: str, xl_name: str) -> str:
    """Markdown rapor (deterministik)."""
    counts = grid_counts(grid)
    items = diff_items(grid)
    total = sum(counts[s] for s in CELL_STATUSES)
    themes = []
    for it in items:
        if it["tema"] not in themes:
            themes.append(it["tema"])
    L = [
        "# Öğretim Programı İnceleme Raporu",
        "",
        f"- PDF: **{pdf_name}**",
        f"- Sistem Excel'i: **{xl_name}**",
        f"- Tarih: {datetime.now():%d.%m.%Y %H:%M}",
        "",
        "## 1. Özet",
        "",
        f"Karşılaştırılan hücre sayısı: **{total}**.",
        "",
        f"- Aynı: {counts[AYNI]}",
        f"- Yalnızca biçim farkı (boşluk, satır sonu tiresi, büyük/küçük harf): {counts[BICIM]}",
        f"- Farklı: **{counts[FARKLI]}**",
        f"- Yalnızca sistemde: **{counts[SADECE_SISTEM]}**",
        f"- Yalnızca PDF'de: **{counts[SADECE_PDF]}**",
        f"- PDF metin katmanından tamamlanan hücre: {counts['tamamlanan_metin_katmani']}",
        f"- Gemini ile bulunup PDF metin katmanında doğrulanan hücre: {counts['tamamlanan_gemini']}",
        f"- Gemini önerisi olup PDF metin katmanında doğrulanamayan hücre: {counts['gemini_dogrulanamayan']}",
        "",
    ]
    if grid.notes:
        L += ["Notlar:", ""] + [f"- {n}" for n in grid.notes] + [""]
    by_field: dict[str, int] = {}
    for it in items:
        by_field[it["alan"]] = by_field.get(it["alan"], 0) + 1
    if by_field:
        L += ["## 2. Alanlara göre farklar", "", "| Alan | Fark sayısı |", "|---|---|"]
        L += [f"| {a} | {n} |" for a, n in sorted(by_field.items(), key=lambda x: -x[1])]
        L.append("")
    L += ["## 3. Tema bazında farklar", ""]
    if not items:
        L.append("Fark bulunmadı.")
    for t in themes:
        L += [f"### {t}", ""]
        for it in (x for x in items if x["tema"] == t):
            where = " · ".join(x for x in (it["oc"], it["bilesen"], it["alan"]) if x)
            L.append(f"- **{where}** — {STATUS_TR.get(it['durum'], it['durum'])}" + (f" ({it['not']})" if it["not"] else ""))
            L.append(f"  - Sistem{f' [{it['excel']}]' if it['excel'] else ''}: {_short(it['sistem']) or '(boş)'}")
            L.append(f"  - PDF{f' [s.{it['sayfa']}]' if it['sayfa'] else ''}: {_short(it['pdf']) or '(boş)'}")
            if it["kaynak"] in (SRC_TEXT_LAYER, SRC_GEMINI):
                L.append(f"  - PDF değerinin kaynağı: {it['kaynak']}")
            if it["gemini_dogrulanamayan"]:
                L.append(f"  - Gemini önerisi (doğrulanamadı, tabloya alınmadı): {_short(it['gemini_dogrulanamayan'])}")
        L.append("")
    L += [
        "## 4. Yöntem",
        "",
        "- Karşılaştırma deterministiktir: temalar öğrenme çıktısı kodlarıyla, öğrenme çıktıları kod kimliğiyle, "
        "süreç bileşenleri sırayla eşleştirilir; metinler boşluk/tire/büyük-küçük harf dışında birebir karşılaştırılır.",
        "- PDF değerleri PDF'nin metin katmanından alınır. Çıkarımda boş kalan hücreler önce sistem metninin PDF metin "
        "katmanında aranmasıyla, isteğe bağlı olarak Gemini'nin sayfa görüntüsünde gösterdiği metnin PDF metin "
        "katmanında doğrulanmasıyla tamamlanır. Doğrulanamayan Gemini önerileri tabloya alınmaz.",
        "",
    ]
    return "\n".join(L)


GEMINI_REPORT_SYSTEM = (
    "Sen öğretim programı inceleme uzmanlarına rapor hazırlayan bir editörsün. Yalnızca sana verilen raporu "
    "daha anlaşılır ve düzenli hale getir. Kurallar: yeni bir olgu, sayı, kod, sayfa numarası veya alıntı "
    "EKLEME; verilen sayıları, kodları ve alıntıları DEĞİŞTİRME; farkları yorumlarken kesin hüküm verme, "
    "incelenmesi gereken noktaları öne çıkar; Türkçe yaz; Markdown kullan (başlıklar, maddeler)."
)


def gemini_report_prompt(markdown_report: str) -> str:
    return (
        "Aşağıdaki deterministik karşılaştırma raporunu inceleme uzmanları için raporlaştır. Başa kısa bir "
        "yönetici özeti, ardından öncelikli inceleme noktaları (ör. yalnızca bir tarafta olan içerik, farklı "
        "metinler, doğrulanamayan Gemini önerileri), sonra tema bazında bulgular koy. Alıntıları olduğu gibi bırak.\n\n"
        "--- RAPOR ---\n" + markdown_report
    )


# ---------------------------------------------------------------- Word


def _shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _inline(par, text: str) -> None:
    """**kalın** işaretli satırı paragrafa yazar."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        run = par.add_run(part)
        run.bold = i % 2 == 1


def _markdown_to_docx(doc: Document, md: str) -> None:
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i + 1].strip()):
            head = [c.strip() for c in ln.strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Table Grid"
            for j, h in enumerate(head):
                t.rows[0].cells[j].text = h.replace("**", "")
            for r in rows:
                cells = t.add_row().cells
                for j, v in enumerate(r[: len(head)]):
                    cells[j].text = v.replace("**", "")
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 3))
        elif re.match(r"^\s{2,}[-*]\s+", ln):
            _inline(doc.add_paragraph(style="List Bullet 2"), re.sub(r"^\s*[-*]\s+", "", ln))
        elif re.match(r"^[-*]\s+", ln):
            _inline(doc.add_paragraph(style="List Bullet"), ln[2:].strip())
        elif re.match(r"^\d+\.\s+", ln):
            _inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", ln))
        else:
            _inline(doc.add_paragraph(), ln)
        i += 1


DIFF_FILL = {FARKLI: "FDE2E1", SADECE_SISTEM: "FFF1D6", SADECE_PDF: "FFF1D6", BICIM: "DEEBF7"}


def build_grid_xlsx(grid: Grid) -> bytes:
    """Nihai karşılaştırma tablosu: her alan için Sistem | PDF | Durum sütunları yan yana."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from excel_export import _xl_safe
    from grid_compare import row_status

    wb = Workbook()
    ws = wb.active
    ws.title = "Nihai karşılaştırma"
    head = ["#", "Satır durumu", "Tema", "Öğrenme çıktısı", "Bileşen"]
    for h in grid.headers:
        head += [f"{h} — Sistem", f"{h} — PDF", f"{h} — Durum"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E7ECF3")
    for r in grid.rows:
        cells = grid.cells_of_row(r, all_levels=True)
        vals = [r.id + 1, STATUS_TR.get(row_status(grid, r), ""), r.theme_label, r.lo_code, r.component]
        for h in grid.headers:
            c = cells.get(h)
            if c is None:
                vals += ["", "", ""]
                continue
            src = f" [{c.pdf_source}]" if c.pdf_source in (SRC_TEXT_LAYER, SRC_GEMINI) else ""
            vals += [c.sys, c.pdf, STATUS_TR.get(c.status, c.status) + src + (f" — {c.note}" if c.note else "")]
        ws.append([_xl_safe(v) for v in vals])
        for j, h in enumerate(grid.headers):
            c = cells.get(h)
            if c is not None and c.status in DIFF_FILL:
                for k in range(3):
                    ws.cell(ws.max_row, 6 + 3 * j + k).fill = PatternFill("solid", fgColor=DIFF_FILL[c.status])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14 if col[0].column <= 5 else 36
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "F2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_docx(grid: Grid, report_md: str, pdf_name: str, xl_name: str, gemini_edited: bool = False) -> bytes:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    if gemini_edited:
        p = doc.add_paragraph()
        r = p.add_run(
            "Bu raporun anlatım bölümü Gemini ile düzenlenmiştir. Kaynak veriler deterministik karşılaştırmadır; "
            "kesin değerler için ekteki fark tablosuna bakınız."
        )
        r.italic = True
        r.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
    _markdown_to_docx(doc, report_md)

    items = diff_items(grid)
    doc.add_page_break()
    doc.add_heading("Ek: Deterministik fark tablosu", level=1)
    doc.add_paragraph(f"PDF: {pdf_name} · Sistem Excel'i: {xl_name} · {len(items)} fark")
    if items:
        cols = ["Tema", "ÖÇ / bileşen", "Alan", "Durum", "Sistem", "PDF", "Not / kaynak"]
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Table Grid"
        for j, h in enumerate(cols):
            c = t.rows[0].cells[j]
            c.text = h
            c.paragraphs[0].runs[0].bold = True
            _shade(c, "E7ECF3")
        for it in items:
            cells = t.add_row().cells
            note = "; ".join(
                x
                for x in (
                    it["not"],
                    it["kaynak"] if it["kaynak"] in (SRC_TEXT_LAYER, SRC_GEMINI) else "",
                    f"Gemini önerisi doğrulanamadı: {_short(it['gemini_dogrulanamayan'], 160)}" if it["gemini_dogrulanamayan"] else "",
                )
                if x
            )
            vals = [
                it["tema"],
                " · ".join(x for x in (it["oc"], it["bilesen"]) if x),
                it["alan"],
                STATUS_TR.get(it["durum"], it["durum"]),
                _short(it["sistem"], 600) + (f" [{it['excel']}]" if it["excel"] else ""),
                _short(it["pdf"], 600) + (f" [s.{it['sayfa']}]" if it["sayfa"] else ""),
                note,
            ]
            for j, v in enumerate(vals):
                cells[j].text = v
                for par in cells[j].paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(8)
            if it["durum"] in DIFF_FILL:
                _shade(cells[3], DIFF_FILL[it["durum"]])
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Pt(40)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
